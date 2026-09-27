"""Small local web console layered on top of the existing data pipeline."""

from __future__ import annotations

import cgi
import json
import mimetypes
import os
import re
import subprocess
import sys
import uuid
from dataclasses import replace
from datetime import date
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from .config import validate_config
from .agent_context import DEFAULT_AGENT_EXPORT_DIR, _history_points, generate_agent_context
from .models import DataPoint
from .pipeline.normalize import normalize_period
from .pipeline.normalized import load_period_points
from .providers.nbs_file import NBSFileProvider
from .report.generator import generate_excel, generate_html
from .storage.sqlite import SQLiteStore


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CONFIG_DIR = PROJECT_ROOT / "config"


def _project_path(value: str | Path) -> Path:
    path = Path(value).expanduser()
    return path if path.is_absolute() else PROJECT_ROOT / path


AGENT_EXPORT_ROOT = _project_path(DEFAULT_AGENT_EXPORT_DIR)


def _periods(data_root: Path) -> list[str]:
    values: set[str] = set()
    for path in (data_root / "normalized").glob("*/*/*.json"):
        if re.fullmatch(r"\d{4}-\d{2}", path.parent.name):
            values.add(path.parent.name)
    return sorted(values, reverse=True)


def _load_points(data_root: Path, period: str) -> list[DataPoint]:
    return load_period_points(data_root, period)


def _previous_period(period: str) -> str:
    year, month = map(int, period.split("-"))
    return f"{year - 1:04d}-12" if month == 1 else f"{year:04d}-{month - 1:02d}"


def _latest_completed_quarter(period: str) -> str:
    year, month = map(int, period.split("-"))
    quarter = (month - 1) // 3
    if quarter == 0:
        return f"{year - 1:04d}-Q4"
    return f"{year:04d}-Q{quarter}"


def _previous_quarter(quarter_period: str) -> str:
    year, quarter = quarter_period.split("-Q")
    year_number, quarter_number = int(year), int(quarter)
    if quarter_number == 1:
        return f"{year_number - 1:04d}-Q4"
    return f"{year_number:04d}-Q{quarter_number - 1}"


def _quarter_key(period: str) -> tuple[int, int]:
    year, quarter = period.split("-Q")
    return int(year), int(quarter)


def _latest_available_quarter_points(data_root: Path, indicator_ids: set[str], maximum: str) -> dict[str, DataPoint]:
    """Return the newest uploaded quarter for each indicator up to maximum."""
    result: dict[str, DataPoint] = {}
    root = data_root / "normalized"
    maximum_key = _quarter_key(maximum)
    for indicator_id in indicator_ids:
        candidates = []
        for directory in (root / indicator_id).glob("????-Q[1-4]"):
            if _quarter_key(directory.name) <= maximum_key:
                candidates.append(directory.name)
        for candidate in sorted(candidates, key=_quarter_key, reverse=True):
            point = next((item for item in load_period_points(data_root, candidate) if item.indicator_id == indicator_id), None)
            if point is not None:
                result[indicator_id] = point
                break
    return result


def _json(handler: BaseHTTPRequestHandler, value: object, status: int = 200) -> None:
    body = json.dumps(value, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def _safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(name).name) or "upload.dat"


def _analysis_payload(row: object) -> dict[str, object]:
    try:
        result = json.loads(str(row["result_json"]))  # type: ignore[index]
    except (TypeError, json.JSONDecodeError):
        result = {}
    try:
        warnings = json.loads(str(row["warnings_json"]))  # type: ignore[index]
    except (TypeError, json.JSONDecodeError):
        warnings = []
    return {
        "analysis_id": int(row["analysis_id"]),  # type: ignore[index]
        "period": row["period"],  # type: ignore[index]
        "package_id": row["package_id"],  # type: ignore[index]
        "version": row["version"],  # type: ignore[index]
        "status": row["status"],  # type: ignore[index]
        "result": result,
        "markdown": row["result_markdown"],  # type: ignore[index]
        "warnings": warnings,
        "created_at": row["created_at"],  # type: ignore[index]
        "confirmed_at": row["confirmed_at"],  # type: ignore[index]
    }


def _analysis_refs(result: object) -> set[str]:
    refs: set[str] = set()
    if isinstance(result, dict):
        for key in ("evidence_refs", "data_refs"):
            evidence = result.get(key)
            if isinstance(evidence, list):
                refs.update(str(item) for item in evidence if isinstance(item, str))
        for value in result.values():
            refs.update(_analysis_refs(value))
    elif isinstance(result, list):
        for value in result:
            refs.update(_analysis_refs(value))
    return refs


def _confirmed_analysis(store: SQLiteStore, period: str) -> dict[str, object] | None:
    row = next((item for item in store.analyses(period=period) if item["status"] == "confirmed"), None)
    if row is None:
        return None
    try:
        result = json.loads(str(row["result_json"]))
    except (TypeError, json.JSONDecodeError):
        return None
    return result if isinstance(result, dict) else None


def _validate_analysis(data_root: Path, period: str, result: object) -> tuple[list[str], str | None]:
    warnings: list[str] = []
    if not isinstance(result, dict):
        return ["分析结果必须是 JSON 对象"], None
    if str(result.get("period", "")) != period:
        return [f"分析期不匹配：需要 {period}"], None
    themes = result.get("themes")
    if not isinstance(themes, dict) or not themes:
        warnings.append("未提供 themes 主题分析")
    else:
        expected_themes = {"生产与供给", "需求与收入", "价格与利润", "货币与信用", "财政与房地产", "外部平衡"}
        missing_themes = sorted(expected_themes - set(themes))
        if missing_themes:
            warnings.append("缺少主题分析：" + "、".join(missing_themes))
        required_fields = ("data_facts", "what_data_indicates", "cross_theme_links", "logic_chain", "future_implications", "watchlist", "data_refs")
        for theme_name, theme in themes.items():
            if not isinstance(theme, dict):
                warnings.append(f"主题 {theme_name} 不是对象")
                continue
            missing_fields = [field for field in required_fields if field not in theme]
            if missing_fields:
                warnings.append(f"主题 {theme_name} 缺少扩展字段：" + ", ".join(missing_fields))
    indicators, _, _ = validate_config(CONFIG_DIR)
    known = {item.indicator_id for item in indicators}
    known.update(point.indicator_id for point in _load_points(data_root, period))
    unknown_refs = sorted(_analysis_refs(result) - known)
    if unknown_refs:
        warnings.append("引用了未在本期数据中找到的指标：" + ", ".join(unknown_refs))
    package_id = str(result.get("package_id") or "") or None
    manifest_path = AGENT_EXPORT_ROOT / period / "manifest.json"
    if package_id and manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            expected = str(manifest.get("package_id") or "")
            if expected and expected != package_id:
                warnings.append("package_id 与当前 Agent 包不一致")
        except (OSError, json.JSONDecodeError):
            warnings.append("当前 Agent 包 manifest 无法读取")
    return warnings, package_id


def _artifact_metadata(row: object) -> dict[str, object]:
    try:
        value = json.loads(str(row["metadata_json"]))  # type: ignore[index]
        return value if isinstance(value, dict) else {}
    except (TypeError, json.JSONDecodeError):
        return {}


def _artifact_vintage(row: object) -> str | None:
    """Read the exact observation vintage from the immutable normalized artifact."""
    metadata = _artifact_metadata(row)
    candidates = [metadata.get("normalized_artifact"), row["path"]]  # type: ignore[index]
    for candidate in candidates:
        if not candidate:
            continue
        try:
            payload = json.loads(Path(str(candidate)).read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError):
            continue
        rows = payload if isinstance(payload, list) else [payload]
        for item in rows:
            if isinstance(item, dict) and item.get("retrieved_at"):
                return str(item["retrieved_at"])
    return None


def _remove_owned_file(path_value: object, data_root: Path) -> bool:
    """Remove only a project-owned file under data/, never a user source file."""
    if not path_value:
        return False
    try:
        path = Path(str(path_value)).resolve()
        root = data_root.resolve()
        if root != path and root not in path.parents:
            return False
        if path.is_file():
            path.unlink()
            return True
    except (OSError, ValueError):
        return False
    return False


def _upload_json(row: object, store: SQLiteStore) -> dict[str, object]:
    metadata = _artifact_metadata(row)
    vintage = _artifact_vintage(row)
    value = None
    if vintage:
        observation = store.connection.execute(
            "SELECT value FROM observations WHERE indicator_id=? AND period=? AND vintage=?",
            (row["indicator_id"], row["period"], vintage),  # type: ignore[index]
        ).fetchone()
        if observation:
            try:
                value = json.loads(observation["value"])
            except (TypeError, json.JSONDecodeError):
                value = observation["value"]
    source_file = str(metadata.get("source_file") or "")
    file_name = Path(source_file).name if source_file else Path(str(row["path"])).name  # type: ignore[index]
    return {
        "artifact_id": int(row["artifact_id"]),  # type: ignore[index]
        "indicator_id": row["indicator_id"],  # type: ignore[index]
        "indicator_name": row["indicator_name"] or row["indicator_id"],  # type: ignore[index]
        "period": row["period"],  # type: ignore[index]
        "source": row["source"],  # type: ignore[index]
        "file_name": file_name,
        "imported_at": row["retrieved_at"],  # type: ignore[index]
        "vintage": vintage,
        "value": value,
        "unit": row["indicator_unit"],  # type: ignore[index]
        "duplicate_count": int(row["duplicate_count"]) if "duplicate_count" in row.keys() else 1,  # type: ignore[index]
        "normalized_file": Path(str(metadata.get("normalized_artifact"))).name if metadata.get("normalized_artifact") else None,
    }


def _delete_artifact(store: SQLiteStore, row: object, data_root: Path) -> dict[str, object]:
    """Remove one project-owned artifact and its exact observation vintage."""
    metadata = _artifact_metadata(row)
    vintage = _artifact_vintage(row)
    raw_path = row["path"]  # type: ignore[index]
    normalized_path = metadata.get("normalized_artifact")
    source_file = metadata.get("source_file")
    result = store.delete_upload(int(row["artifact_id"]), vintage=vintage)  # type: ignore[index]
    # A digest may be shared by multiple imports. Only remove a project file
    # after the last raw_artifacts row referencing it has been deleted.
    if store.connection.execute("SELECT 1 FROM raw_artifacts WHERE path=? LIMIT 1", (raw_path,)).fetchone() is None:
        _remove_owned_file(raw_path, data_root)
    if normalized_path and store.connection.execute(
        "SELECT 1 FROM raw_artifacts WHERE json_extract(metadata_json, '$.normalized_artifact')=? LIMIT 1",
        (str(normalized_path),),
    ).fetchone() is None:
        _remove_owned_file(normalized_path, data_root)
    if source_file and store.connection.execute(
        "SELECT 1 FROM raw_artifacts WHERE json_extract(metadata_json, '$.source_file')=? LIMIT 1",
        (str(source_file),),
    ).fetchone() is None:
        _remove_owned_file(source_file, data_root)
    return {**(result or {}), "warning": None if vintage else "未能定位该记录的标准化版本，未删除观测值"}


def _xlsx_rows(path: Path) -> list[dict[str, object]]:
    """Read simple official workbooks without changing the existing file provider."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Excel 导入需要安装 openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    period_names = {"period", "date", "time", "dt", "月份", "日期", "统计期", "时间"}
    value_names = {"value", "data", "数值", "指标值", "同比", "环比", "增速", "当月值"}
    rows: list[dict[str, object]] = []
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        for header_index, header in enumerate(values[:12]):
            names = [str(item).strip().lower() if item is not None else "" for item in header]
            period_col = next((i for i, name in enumerate(names) if name in period_names), None)
            value_col = next((i for i, name in enumerate(names) if name in value_names), None)
            if period_col is None or value_col is None:
                continue
            for row in values[header_index + 1:]:
                if not row or max(period_col, value_col) >= len(row):
                    continue
                if row[period_col] is not None and row[value_col] is not None:
                    try:
                        canonical_period = normalize_period(row[period_col], "monthly")
                    except (ValueError, TypeError):
                        # Ignore footnotes and source rows below official tables.
                        continue
                    if isinstance(row[value_col], (int, float)) and not isinstance(row[value_col], bool):
                        rows.append({"period": canonical_period, "value": row[value_col], "unit": "%"})
            break
    return rows


class WebHandler(BaseHTTPRequestHandler):
    data_root = _project_path(os.environ.get("MACRO_DATA_ROOT", str(PROJECT_ROOT / "data")))
    reports_root = _project_path(os.environ.get("MACRO_REPORT_DIR", str(PROJECT_ROOT / "reports")))
    batch_cache: dict[str, dict[str, object]] = {}

    # The monthly upload package follows the public 01-42 checklist.  The
    # first version of the web page only listed 01-20, which silently skipped
    # the external, rates, international, and fiscal workbooks.
    BATCH_INDICATORS = {
        1: ("industrial_value_added_yoy", "nbs"),
        2: ("retail_sales_yoy", "nbs"),
        3: ("fixed_asset_investment_ytd_yoy", "nbs"),
        4: ("cpi_yoy", "nbs"),
        5: ("cpi_core_yoy", "nbs"),
        6: ("ppi_yoy", "nbs"),
        7: ("pmi_manufacturing", "nbs"),
        8: ("m1_yoy", "pboc"),
        9: ("m0_yoy", "pboc"),
        10: ("m2_yoy", "pboc"),
        11: ("social_financing_stock_yoy", "pboc"),
        12: ("social_financing_increment_monthly", "pboc"),
        13: ("rmb_loans_stock_yoy", "pboc"),
        14: ("rmb_loans_increment_monthly", "pboc"),
        15: ("rmb_deposits_stock_yoy", "pboc"),
        16: ("loan_prime_rate_1y", "pboc"),
        17: ("industrial_profit_yoy", "nbs"),
        18: ("electricity_generation_yoy", "nbs"),
        19: ("finished_goods_inventory_yoy", "nbs"),
        20: ("capacity_utilization_rate", "nbs"),
        21: ("real_estate_sales_area_ytd_yoy", "nbs"),
        22: ("exports_rmb_monthly", "nbs"),
        23: ("imports_rmb_monthly", "nbs"),
        24: ("trade_balance_rmb_monthly", "nbs"),
        25: ("bank_cross_border_receipts_payments_balance_usd", "safe"),
        26: ("foreign_exchange_reserves_usd", "safe"),
        27: ("cny_usd_month_end", "chinamoney"),
        28: ("bank_settlement_balance_usd", "safe"),
        29: ("international_balance_current_account", "safe"),
        30: ("china_10y_gov_yield", "chinabond"),
        31: ("cny_usd_month_average", "rates"),
        32: ("china_10y_gov_yield_month_end", "rates"),
        33: ("dr007_month_average", "rates"),
        34: ("repo_7d_month_average", "rates"),
        35: ("us_fed_funds_rate", "fred"),
        36: ("us_10y_treasury_yield", "fred"),
        37: ("dollar_index", "fred"),
        38: ("general_public_budget_revenue_ytd_yoy", "mof"),
        39: ("general_public_budget_expenditure_ytd_yoy", "mof"),
        40: ("central_general_budget_revenue_ytd_yoy", "mof"),
        41: ("local_general_budget_revenue_ytd_yoy", "mof"),
        42: ("government_bond_net_financing_monthly", "mof"),
    }

    def log_message(self, fmt: str, *args: object) -> None:
        # Keep the console useful without ever printing request bodies or secrets.
        print(f"[web] {self.address_string()} {fmt % args}")

    def do_GET(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path in {"/workbench", "/workbench/"}:
            self._serve(PROJECT_ROOT / "web" / "workbench.html")
            return
        if parsed.path == "/api/periods":
            _json(self, {"periods": _periods(self.data_root)})
            return
        if parsed.path == "/api/dashboard":
            period = parse_qs(parsed.query).get("period", [""])[0]
            if not re.fullmatch(r"\d{4}-\d{2}", period):
                _json(self, {"error": "period must be YYYY-MM"}, 400)
                return
            _json(self, self._dashboard(period))
            return
        if parsed.path == "/api/history":
            indicator_id = parse_qs(parsed.query).get("indicator", [""])[0]
            history = []
            for period in _periods(self.data_root):
                for point in _load_points(self.data_root, period):
                    if point.indicator_id == indicator_id:
                        history.append({"period": point.period, "value": point.value, "unit": point.unit})
            _json(self, {"indicator_id": indicator_id, "history": sorted(history, key=lambda row: row["period"])})
            return
        if parsed.path == "/api/series":
            indicators, _, _ = validate_config(CONFIG_DIR)
            requested = {item.indicator_id for item in indicators}
            series: dict[str, list[dict[str, object]]] = {item.indicator_id: [] for item in indicators}
            for period in _periods(self.data_root):
                for point in _load_points(self.data_root, period):
                    if point.indicator_id in requested and point.value is not None:
                        series[point.indicator_id].append({"period": point.period, "value": point.value, "unit": point.unit})
            for values in series.values():
                values.sort(key=lambda row: str(row["period"]))
            _json(self, {"series": series, "indicators": [{"indicator_id": item.indicator_id, "name": item.name, "theme": item.theme, "unit": item.unit, "frequency": item.frequency, "value_type": item.value_type, "definition": item.definition} for item in indicators]})
            return
        if parsed.path == "/api/uploads":
            period = parse_qs(parsed.query).get("period", [""])[0]
            if period and not re.fullmatch(r"\d{4}-\d{2}", period):
                _json(self, {"error": "period must be YYYY-MM"}, 400)
                return
            store = SQLiteStore(self.data_root / "macro_observer.sqlite")
            try:
                rows = [_upload_json(row, store) for row in store.uploads(period=period or None)]
            finally:
                store.close()
            _json(self, {"period": period or None, "uploads": rows, "count": len(rows)})
            return
        if parsed.path == "/api/agent-context":
            period = parse_qs(parsed.query).get("period", [""])[0]
            if not re.fullmatch(r"\d{4}-\d{2}", period):
                _json(self, {"error": "period must be YYYY-MM"}, 400)
                return
            indicators, _, _ = validate_config(CONFIG_DIR)
            _json(self, self._agent_context_payload(period, indicators))
            return
        if parsed.path.startswith("/agent-pack/"):
            match = re.fullmatch(r"/agent-pack/(\d{4}-\d{2})\.zip", parsed.path)
            if not match:
                _json(self, {"error": "not found"}, 404)
                return
            self._serve_external(AGENT_EXPORT_ROOT / f"{match.group(1)}.zip", "application/zip")
            return
        if parsed.path == "/api/analysis":
            period = parse_qs(parsed.query).get("period", [""])[0]
            if period and not re.fullmatch(r"\d{4}-\d{2}", period):
                _json(self, {"error": "period must be YYYY-MM"}, 400)
                return
            store = SQLiteStore(self.data_root / "macro_observer.sqlite")
            try:
                rows = [_analysis_payload(row) for row in store.analyses(period=period or None)]
            finally:
                store.close()
            _json(self, {"period": period or None, "analyses": rows, "confirmed": next((row for row in rows if row["status"] == "confirmed"), None)})
            return
        if parsed.path.startswith("/reports/"):
            self._serve(PROJECT_ROOT / parsed.path.lstrip("/"))
            return
        if parsed.path.startswith("/data/agent_context/"):
            self._serve(PROJECT_ROOT / parsed.path.lstrip("/"))
            return
        self._serve(PROJECT_ROOT / "web" / ("index.html" if parsed.path in {"", "/"} else parsed.path.lstrip("/")))

    def _serve(self, path: Path) -> None:
        try:
            resolved = path.resolve()
            if not resolved.is_file() or PROJECT_ROOT not in resolved.parents:
                raise FileNotFoundError
            body = resolved.read_bytes()
        except (OSError, FileNotFoundError):
            _json(self, {"error": "not found"}, 404)
            return
        self.send_response(200)
        self.send_header("Content-Type", mimetypes.guess_type(str(path))[0] or "application/octet-stream")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _serve_external(self, path: Path, content_type: str) -> None:
        """Serve only the generated Agent archive outside the project root."""
        try:
            body = path.resolve().read_bytes()
        except OSError:
            _json(self, {"error": "Agent 包不存在，请先导出"}, 404)
            return
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f'attachment; filename="{path.name}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _agent_context_payload(self, period: str, indicators: list[object]) -> dict[str, object]:
        definitions = list(indicators)
        current = _load_points(self.data_root, period)
        quarter_ids = {item.indicator_id for item in definitions if item.frequency == "quarterly"}
        quarter_period = _latest_completed_quarter(period)
        current_ids = {point.indicator_id for point in current}
        quarter_points = [point for indicator_id, point in _latest_available_quarter_points(self.data_root, quarter_ids, quarter_period).items() if indicator_id not in current_ids]
        target = generate_agent_context(period, [*current, *quarter_points], definitions, output_dir=AGENT_EXPORT_ROOT, history_points=_history_points(self.data_root, period))
        return {
            "period": period,
            "path": str(target),
            "prompt": str(target / "analysis-prompt.md"),
            "zip": f"/agent-pack/{period}.zip",
            "download": f"/agent-pack/{period}.zip",
            "output_dir": str(AGENT_EXPORT_ROOT),
        }

    def _dashboard(self, period: str) -> dict[str, object]:
        indicators, sources, _ = validate_config(CONFIG_DIR)
        quarter_ids = {item.indicator_id for item in indicators if item.frequency == "quarterly"}
        quarter_period = _latest_completed_quarter(period)
        quarter_points = _latest_available_quarter_points(self.data_root, quarter_ids, quarter_period)
        previous_quarter_points = _latest_available_quarter_points(self.data_root, quarter_ids, _previous_quarter(quarter_period))
        current = {point.indicator_id: point for point in _load_points(self.data_root, period)}
        previous = {point.indicator_id: point for point in _load_points(self.data_root, _previous_period(period))}
        current.update(quarter_points)
        previous.update(previous_quarter_points)
        rows = []
        for indicator in indicators:
            point = current.get(indicator.indicator_id)
            old = previous.get(indicator.indicator_id)
            value = point.value if point else None
            prior = old.value if old else None
            change = value - prior if isinstance(value, (int, float)) and isinstance(prior, (int, float)) else None
            rows.append({"indicator_id": indicator.indicator_id, "name": indicator.name, "theme": indicator.theme, "unit": indicator.unit, "value": value, "previous": prior, "change": change, "source": point.source if point else indicator.primary_source, "source_url": point.source_url if point else None, "homepage_core": indicator.homepage_core, "status": "已更新" if point else "待接入"})
        available = sum(row["value"] is not None for row in rows)
        source_aliases = {"nbs_file": "nbs", "web_upload": "nbs"}
        source_states = [{"name": source_name, "label": config.get("name", source_name), "status": "已导入" if any(source_aliases.get(str(row["source"]), str(row["source"])) == source_name and row["value"] is not None for row in rows) else "待上传"} for source_name, config in sources.items()]
        package_monthly = [
            item for number, item in self.BATCH_INDICATORS.items()
            if next((definition.frequency for definition in indicators if definition.indicator_id == item[0]), "monthly") == "monthly"
        ]
        monthly_available = sum(current.get(indicator_id) is not None for indicator_id, _ in package_monthly)
        quarter_available = len(quarter_points)
        store = SQLiteStore(self.data_root / "macro_observer.sqlite")
        confirmed = next((row for row in store.analyses(period=period) if row["status"] == "confirmed"), None)
        store.close()
        all_available = available
        all_total = len(rows)
        return {"period": period, "previous_period": _previous_period(period), "updated_at": date.today().isoformat(), "available": all_available, "total": all_total, "rows": rows, "source_states": source_states, "package_status": {"monthly_available": monthly_available, "monthly_total": len(package_monthly), "quarterly_available": int(quarter_available), "quarterly_total": len(quarter_ids), "all_available": all_available, "all_total": all_total, "quarter_period": quarter_period}, "analysis_status": confirmed["status"] if confirmed else "missing"}

    def do_POST(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path == "/api/report":
            payload = self._body_json()
            period = str(payload.get("period", ""))
            if not re.fullmatch(r"\d{4}-\d{2}", period):
                _json(self, {"error": "period must be YYYY-MM"}, 400)
                return
            points = _load_points(self.data_root, period)
            store = SQLiteStore(self.data_root / "macro_observer.sqlite")
            analysis = _confirmed_analysis(store, period)
            store.close()
            html_path = generate_html(period, points, analysis=analysis, output_dir=self.reports_root)
            excel_path = generate_excel(period, points, output_dir=self.reports_root)
            _json(self, {"html": f"/reports/{period}/monthly-report.html", "excel": str(excel_path.relative_to(PROJECT_ROOT)).replace("\\", "/")})
            return
        if parsed.path == "/api/agent-context":
            payload = self._body_json()
            period = str(payload.get("period", ""))
            if not re.fullmatch(r"\d{4}-\d{2}", period):
                _json(self, {"error": "period must be YYYY-MM"}, 400)
                return
            indicators, _, _ = validate_config(CONFIG_DIR)
            _json(self, self._agent_context_payload(period, indicators))
            return
        if parsed.path == "/api/analysis/import":
            self._import_analysis()
            return
        if parsed.path == "/api/analysis/confirm":
            payload = self._body_json()
            try:
                analysis_id = int(payload.get("analysis_id", 0))
            except (TypeError, ValueError):
                analysis_id = 0
            store = SQLiteStore(self.data_root / "macro_observer.sqlite")
            row = store.confirm_analysis(analysis_id) if analysis_id > 0 else None
            store.close()
            if row is None:
                _json(self, {"error": "分析结果不存在"}, 404)
            else:
                _json(self, {"status": "confirmed", "analysis": _analysis_payload(row)})
            return
        if parsed.path == "/api/refresh":
            payload = self._body_json()
            period = str(payload.get("period", ""))
            if not re.fullmatch(r"\d{4}-\d{2}", period):
                _json(self, {"error": "period must be YYYY-MM"}, 400)
                return
            log_dir = self.data_root / "logs"
            log_dir.mkdir(parents=True, exist_ok=True)
            log = (log_dir / f"run-{period}.log").open("a", encoding="utf-8")
            subprocess.Popen([sys.executable, "-m", "macro_observer", "run", "--period", period], cwd=PROJECT_ROOT, stdout=log, stderr=subprocess.STDOUT)
            _json(self, {"started": True, "message": "已启动后台整理，完成后刷新页面查看"})
            return
        if parsed.path == "/api/upload":
            self._upload()
            return
        if parsed.path == "/api/uploads/delete":
            self._delete_upload()
            return
        if parsed.path == "/api/uploads/deduplicate":
            self._deduplicate_uploads()
            return
        if parsed.path == "/api/batch-preview":
            self._batch_preview()
            return
        if parsed.path == "/api/batch-commit":
            self._batch_commit()
            return
        _json(self, {"error": "not found"}, 404)

    def _body_json(self) -> dict[str, object]:
        length = int(self.headers.get("Content-Length", "0"))
        try:
            value = json.loads(self.rfile.read(length) or b"{}")
            return value if isinstance(value, dict) else {}
        except (json.JSONDecodeError, ValueError):
            return {}

    def _upload(self) -> None:
        form = cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={"REQUEST_METHOD": "POST", "CONTENT_TYPE": self.headers.get("Content-Type", ""), "CONTENT_LENGTH": self.headers.get("Content-Length", "0")})
        period = str(form.getfirst("period", ""))
        indicator_id = str(form.getfirst("indicator", ""))
        uploaded = form["file"] if "file" in form else None
        if not re.fullmatch(r"\d{4}-\d{2}", period) or not indicator_id or uploaded is None or not getattr(uploaded, "filename", None):
            _json(self, {"error": "需要统计期、指标和文件"}, 400)
            return
        indicators, _, _ = validate_config(CONFIG_DIR)
        indicator = next((item for item in indicators if item.indicator_id == indicator_id), None)
        if indicator is None:
            _json(self, {"error": "未知指标"}, 400)
            return
        incoming = self.data_root / "incoming" / "web"
        incoming.mkdir(parents=True, exist_ok=True)
        original_name = str(uploaded.filename)
        suffix = Path(original_name).suffix.lower()
        content_type = str(getattr(uploaded, "type", "") or "").lower()
        uploaded_bytes = uploaded.file.read()
        if suffix not in {".csv", ".json", ".xlsx", ".xlsm"}:
            if "spreadsheet" in content_type or "excel" in content_type:
                suffix = ".xlsx"
            elif "json" in content_type:
                suffix = ".json"
            elif "csv" in content_type:
                suffix = ".csv"
            elif uploaded_bytes[:4] == b"PK\x03\x04":
                # XLSX/XLSM files are ZIP containers; use the safe parser even
                # when a multipart client mangles a non-ASCII filename.
                suffix = ".xlsx"
        safe_stem = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(original_name).stem).strip("._") or "upload"
        target = incoming / f"{safe_stem}{suffix}"
        target.write_bytes(uploaded_bytes)
        import_path = target
        if target.suffix.lower() in {".xlsx", ".xlsm"}:
            rows = _xlsx_rows(target)
            import_path = target.with_suffix(".json")
            import_path.write_text(json.dumps(rows, ensure_ascii=False, default=str), encoding="utf-8")
        result = NBSFileProvider(data_root=self.data_root).fetch(indicator, period, import_path)
        store = SQLiteStore(self.data_root / "macro_observer.sqlite")
        if result.status == "success":
            store.upsert_definition(indicator, updated_at=date.today().isoformat())
            store.insert_fetch_result(result, provider="web_upload")
        store.close()
        _json(self, {"status": result.status, "errors": result.errors, "points": len(result.points), "file": str(target.relative_to(PROJECT_ROOT)).replace("\\", "/")})

    def _import_analysis(self) -> None:
        form = cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={"REQUEST_METHOD": "POST", "CONTENT_TYPE": self.headers.get("Content-Type", ""), "CONTENT_LENGTH": self.headers.get("Content-Length", "0")})
        period = str(form.getfirst("period", ""))
        uploaded = form["file"] if "file" in form else None
        if not re.fullmatch(r"\d{4}-\d{2}", period) or uploaded is None or not getattr(uploaded, "filename", None):
            _json(self, {"error": "需要统计期和 analysis-result.json 文件"}, 400)
            return
        try:
            result = json.loads(uploaded.file.read().decode("utf-8-sig"))
        except (UnicodeError, json.JSONDecodeError) as exc:
            _json(self, {"error": f"分析结果 JSON 无法解析：{exc}"}, 400)
            return
        warnings, package_id = _validate_analysis(self.data_root, period, result)
        markdown = str(form.getfirst("markdown", "")) or None
        store = SQLiteStore(self.data_root / "macro_observer.sqlite")
        analysis_id = store.save_analysis(period=period, package_id=package_id, result=result, markdown=markdown, warnings=warnings)
        row = store.analysis(analysis_id)
        store.close()
        _json(self, {"status": "draft", "analysis": _analysis_payload(row), "message": "分析结果已导入，请预览并确认"})

    def _delete_upload(self) -> None:
        payload = self._body_json()
        try:
            artifact_id = int(payload.get("artifact_id", 0))
        except (TypeError, ValueError):
            artifact_id = 0
        if artifact_id <= 0:
            _json(self, {"error": "artifact_id must be a positive integer"}, 400)
            return
        store = SQLiteStore(self.data_root / "macro_observer.sqlite")
        row = store.upload(artifact_id)
        if row is None:
            store.close()
            _json(self, {"error": "上传记录不存在"}, 404)
            return
        result = _delete_artifact(store, row, self.data_root)
        store.close()
        _json(self, {"status": "success", "deleted": result})

    def _deduplicate_uploads(self) -> None:
        payload = self._body_json()
        period = str(payload.get("period", ""))
        if not re.fullmatch(r"\d{4}-\d{2}", period):
            _json(self, {"error": "period must be YYYY-MM"}, 400)
            return
        store = SQLiteStore(self.data_root / "macro_observer.sqlite")
        rows = store.duplicate_uploads(period=period)
        deleted: list[dict[str, object]] = []
        warnings: list[int] = []
        remaining = 0
        try:
            for row in rows:
                result = _delete_artifact(store, row, self.data_root)
                if result.get("warning"):
                    warnings.append(int(row["artifact_id"]))
                deleted.append(result)
            remaining = len(store.uploads(period=period))
        finally:
            store.close()
        _json(self, {"status": "success", "period": period, "kept": remaining, "removed": len(deleted), "warnings": warnings, "deleted": deleted})

    @staticmethod
    def _form_files(form: cgi.FieldStorage) -> list[cgi.FieldStorage]:
        value = form["files"] if "files" in form else []
        if not isinstance(value, list):
            value = [value]
        return [item for item in value if isinstance(item, cgi.FieldStorage) and getattr(item, "filename", None)]

    def _batch_prepare(self, period: str, files: list[cgi.FieldStorage]) -> tuple[str, dict[str, object]]:
        indicators, _, _ = validate_config(CONFIG_DIR)
        definitions = {item.indicator_id: item for item in indicators}
        batch_id = uuid.uuid4().hex
        staging = self.data_root / "incoming" / "web" / "batches" / batch_id
        staging.mkdir(parents=True, exist_ok=True)
        prepared: list[dict[str, object]] = []
        seen_numbers: set[int] = set()
        for uploaded in files:
            name = str(uploaded.filename)
            match = re.match(r"\s*(\d{1,2})[_-]", Path(name).name)
            if not match:
                continue
            number = int(match.group(1))
            if number not in self.BATCH_INDICATORS or number in seen_numbers:
                continue
            seen_numbers.add(number)
            indicator_id, source = self.BATCH_INDICATORS[number]
            indicator = definitions[indicator_id]
            raw = uploaded.file.read()
            suffix = Path(name).suffix.lower()
            if suffix not in {".xlsx", ".xlsm"} and raw[:4] == b"PK\x03\x04":
                suffix = ".xlsx"
            target = staging / f"{number:02d}{suffix or '.xlsx'}"
            target.write_bytes(raw)
            if suffix not in {".xlsx", ".xlsm"}:
                prepared.append({"number": number, "indicator_id": indicator_id, "status": "格式不支持", "message": "月度数据包目前需要 Excel 文件"})
                continue
            rows = _xlsx_rows(target)
            candidates = [row for row in rows if row.get("period") == period and row.get("value") is not None]
            fetch_period = period
            if indicator.frequency == "quarterly":
                candidates = [row for row in rows if row.get("value") is not None and str(row.get("period", "")) <= period]
                if candidates:
                    latest = candidates[-1]
                    raw_period = str(latest["period"])
                    if re.fullmatch(r"\d{4}-Q[1-4]", raw_period):
                        fetch_period = raw_period
                    else:
                        year, month = map(int, raw_period.split("-"))
                        fetch_period = f"{year:04d}-Q{(month - 1) // 3 + 1}"
                    candidates = [{**latest, "period": fetch_period}]
            if not candidates:
                prepared.append({"number": number, "indicator_id": indicator_id, "name": indicator.name, "status": "未找到当期值", "message": "文件中没有可匹配的当期数值"})
                continue
            json_path = staging / f"{number:02d}.json"
            candidates = [{**candidate, "unit": indicator.unit} for candidate in candidates]
            json_path.write_text(json.dumps(candidates, ensure_ascii=False), encoding="utf-8")
            prepared.append({"number": number, "indicator_id": indicator_id, "name": indicator.name, "source": source, "status": "待导入", "period": fetch_period, "value": candidates[0]["value"], "unit": indicator.unit, "path": str(json_path)})
        missing = sorted(set(self.BATCH_INDICATORS) - seen_numbers)
        payload = {"batch_id": batch_id, "period": period, "files": prepared, "recognized": len(seen_numbers), "expected": len(self.BATCH_INDICATORS), "missing_numbers": missing, "staging": str(staging)}
        self.batch_cache[batch_id] = payload
        return batch_id, payload

    def _batch_preview(self) -> None:
        form = cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={"REQUEST_METHOD": "POST", "CONTENT_TYPE": self.headers.get("Content-Type", ""), "CONTENT_LENGTH": self.headers.get("Content-Length", "0")})
        period = str(form.getfirst("period", ""))
        files = self._form_files(form)
        if not re.fullmatch(r"\d{4}-\d{2}", period) or not files:
            _json(self, {"error": "需要统计期和至少一个 Excel 文件"}, 400)
            return
        _, payload = self._batch_prepare(period, files)
        _json(self, {key: value for key, value in payload.items() if key != "staging"})

    def _batch_commit(self) -> None:
        payload = self._body_json()
        batch_id = str(payload.get("batch_id", ""))
        batch = self.batch_cache.get(batch_id)
        if not batch:
            _json(self, {"error": "批次不存在或网页服务已重启，请重新选择文件夹"}, 404)
            return
        indicators, _, _ = validate_config(CONFIG_DIR)
        definitions = {item.indicator_id: item for item in indicators}
        store = SQLiteStore(self.data_root / "macro_observer.sqlite")
        results = []
        for item in batch["files"]:
            if item.get("status") != "待导入":
                results.append(item)
                continue
            indicator = definitions[str(item["indicator_id"])]
            result = NBSFileProvider(data_root=self.data_root).fetch(indicator, str(item["period"]), str(item["path"]))
            source = str(item.get("source", "nbs"))
            if result.status == "success":
                result.points = [replace(point, source=source) for point in result.points]
                store.upsert_definition(indicator, updated_at=date.today().isoformat())
                store.insert_fetch_result(result, provider=source)
            results.append({"number": item["number"], "indicator_id": item["indicator_id"], "name": item.get("name"), "status": result.status, "points": len(result.points), "errors": result.errors})
        store.close()
        self.batch_cache.pop(batch_id, None)
        _json(self, {"period": batch["period"], "results": results, "success": sum(1 for item in results if item.get("status") == "success")})


def serve(host: str = "127.0.0.1", port: int = 8765, *, data_root: Path | str = Path(os.environ.get("MACRO_DATA_ROOT", str(PROJECT_ROOT / "data"))).expanduser()) -> None:
    WebHandler.data_root = _project_path(data_root).resolve()
    WebHandler.reports_root = _project_path(os.environ.get("MACRO_REPORT_DIR", str(PROJECT_ROOT / "reports"))).resolve()
    server = ThreadingHTTPServer((host, port), WebHandler)
    print(f"本地网页：http://{host}:{port}/")
    server.serve_forever()
